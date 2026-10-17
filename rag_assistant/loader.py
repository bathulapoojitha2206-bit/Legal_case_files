"""
Document ingestion and extraction module for Legal Case Files.
Supports PDF, TXT, and DOCX documents with robust parsing and error handling.
"""

import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Dict, List, Optional, Tuple, Union

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


@dataclass
class LoadedDocument:
    """Represents the outcome of a document load operation."""
    filename: str
    documents: List[Document]
    success: bool
    error_message: Optional[str] = None
    warning_message: Optional[str] = None


class DocumentLoaderError(Exception):
    """Custom exception raised for unrecoverable loader errors."""
    pass


class DocumentLoader:
    """
    Ingests and parses legal documents (.pdf, .txt, .docx)
    into LangChain Document objects with detailed metadata.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}

    def __init__(self):
        pass

    def load_file(self, file_path: Union[str, Path]) -> LoadedDocument:
        """
        Load a single file from the filesystem.

        Args:
            file_path: Path to the target file.

        Returns:
            LoadedDocument containing parsed LangChain Documents or error details.
        """
        path = Path(file_path).resolve()
        filename = path.name

        if not path.exists():
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=f"File not found: {path}"
            )

        if not path.is_file():
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=f"Path is not a regular file: {path}"
            )

        file_size = path.stat().st_size
        if file_size == 0:
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                warning_message=f"File is completely empty (0 bytes): {filename}"
            )

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=(
                    f"Unsupported document format '{ext}' for file '{filename}'. "
                    f"Supported formats: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
                )
            )

        try:
            if ext == ".pdf":
                docs = self._load_pdf(path)
            elif ext == ".txt":
                docs = self._load_txt(path)
            elif ext == ".docx":
                docs = self._load_docx(path)
            else:
                docs = []

            if not docs or all(len(d.page_content.strip()) == 0 for d in docs):
                return LoadedDocument(
                    filename=filename,
                    documents=[],
                    success=False,
                    warning_message=f"No readable text could be extracted from: {filename}"
                )

            return LoadedDocument(filename=filename, documents=docs, success=True)

        except Exception as exc:
            logger.error("Failed to parse file %s: %s", filename, exc, exc_info=True)
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=f"Error parsing {filename}: {str(exc)}"
            )

    def load_bytes(self, file_bytes: bytes, filename: str) -> LoadedDocument:
        """
        Load a document from in-memory bytes (useful for Streamlit uploads).

        Args:
            file_bytes: Raw bytes of the file.
            filename: Original file name.

        Returns:
            LoadedDocument object.
        """
        if not file_bytes or len(file_bytes) == 0:
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                warning_message=f"Uploaded file is empty (0 bytes): {filename}"
            )

        ext = Path(filename).suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=(
                    f"Unsupported file format '{ext}' for file '{filename}'. "
                    f"Supported formats: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
                )
            )

        try:
            import io
            stream = io.BytesIO(file_bytes)

            if ext == ".pdf":
                docs = self._load_pdf_stream(stream, filename)
            elif ext == ".txt":
                docs = self._load_txt_bytes(file_bytes, filename)
            elif ext == ".docx":
                docs = self._load_docx_stream(stream, filename)
            else:
                docs = []

            if not docs or all(len(d.page_content.strip()) == 0 for d in docs):
                return LoadedDocument(
                    filename=filename,
                    documents=[],
                    success=False,
                    warning_message=f"No readable text extracted from: {filename}"
                )

            return LoadedDocument(filename=filename, documents=docs, success=True)

        except Exception as exc:
            logger.error("Failed to parse uploaded file %s: %s", filename, exc, exc_info=True)
            return LoadedDocument(
                filename=filename,
                documents=[],
                success=False,
                error_message=f"Parsing error for {filename}: {str(exc)}"
            )

    def load_directory(
        self, directory_path: Union[str, Path], recursive: bool = True
    ) -> Tuple[List[Document], List[LoadedDocument]]:
        """
        Scan a directory and load all supported documents.

        Args:
            directory_path: Directory path to scan.
            recursive: Whether to search subdirectories recursively.

        Returns:
            Tuple of (all_extracted_documents, list_of_load_results_per_file).
        """
        dir_p = Path(directory_path).resolve()
        if not dir_p.exists() or not dir_p.is_dir():
            raise DocumentLoaderError(f"Directory not found or not a valid directory: {dir_p}")

        glob_pattern = "**/*" if recursive else "*"
        all_docs: List[Document] = []
        load_results: List[LoadedDocument] = []

        for item in sorted(dir_p.glob(glob_pattern)):
            if item.is_file():
                ext = item.suffix.lower()
                if ext in self.SUPPORTED_EXTENSIONS:
                    result = self.load_file(item)
                    load_results.append(result)
                    if result.success and result.documents:
                        all_docs.extend(result.documents)

        return all_docs, load_results

    # -------------------------------------------------------------------------
    # Internal Format Parsers
    # -------------------------------------------------------------------------

    def _load_pdf(self, path: Path) -> List[Document]:
        """Extract text page-by-page from a PDF file using pypdf."""
        import pypdf

        documents = []
        filename = path.name

        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            total_pages = len(reader.pages)

            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                page_text = page_text.strip()
                if page_text:
                    documents.append(
                        Document(
                            page_content=page_text,
                            metadata={
                                "source": str(path),
                                "filename": filename,
                                "page": page_idx + 1,
                                "total_pages": total_pages,
                                "file_type": "pdf",
                                "char_count": len(page_text),
                            },
                        )
                    )

        return documents

    def _load_pdf_stream(self, stream: BinaryIO, filename: str) -> List[Document]:
        """Extract text page-by-page from an in-memory PDF stream."""
        import pypdf

        documents = []
        reader = pypdf.PdfReader(stream)
        total_pages = len(reader.pages)

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_text = page_text.strip()
            if page_text:
                documents.append(
                    Document(
                        page_content=page_text,
                        metadata={
                            "source": filename,
                            "filename": filename,
                            "page": page_idx + 1,
                            "total_pages": total_pages,
                            "file_type": "pdf",
                            "char_count": len(page_text),
                        },
                    )
                )

        return documents

    def _load_txt(self, path: Path) -> List[Document]:
        """Load text file with multi-encoding fallback."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        content = None
        last_error = None

        for enc in encodings:
            try:
                with open(path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, LookupError) as e:
                last_error = e
                continue

        if content is None:
            raise DocumentLoaderError(f"Could not decode text file {path.name}: {last_error}")

        content_clean = content.strip()
        if not content_clean:
            return []

        return [
            Document(
                page_content=content_clean,
                metadata={
                    "source": str(path),
                    "filename": path.name,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "txt",
                    "char_count": len(content_clean),
                },
            )
        ]

    def _load_txt_bytes(self, data: bytes, filename: str) -> List[Document]:
        """Decode in-memory text bytes with multi-encoding fallback."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        text = None
        last_error = None

        for enc in encodings:
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError as e:
                last_error = e
                continue

        if text is None:
            raise DocumentLoaderError(f"Could not decode bytes for {filename}: {last_error}")

        clean = text.strip()
        if not clean:
            return []

        return [
            Document(
                page_content=clean,
                metadata={
                    "source": filename,
                    "filename": filename,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "txt",
                    "char_count": len(clean),
                },
            )
        ]

    def _load_docx(self, path: Path) -> List[Document]:
        """Extract text from a DOCX file using python-docx."""
        import docx

        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        full_text = "\n\n".join(paragraphs).strip()

        if not full_text:
            return []

        return [
            Document(
                page_content=full_text,
                metadata={
                    "source": str(path),
                    "filename": path.name,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "docx",
                    "char_count": len(full_text),
                },
            )
        ]

    def _load_docx_stream(self, stream: BinaryIO, filename: str) -> List[Document]:
        """Extract text from an in-memory DOCX stream."""
        import docx

        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        full_text = "\n\n".join(paragraphs).strip()

        if not full_text:
            return []

        return [
            Document(
                page_content=full_text,
                metadata={
                    "source": filename,
                    "filename": filename,
                    "page": 1,
                    "total_pages": 1,
                    "file_type": "docx",
                    "char_count": len(full_text),
                },
            )
        ]
